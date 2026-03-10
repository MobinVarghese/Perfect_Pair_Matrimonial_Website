"""
Export Utilities Module
Helper functions for exporting data to PDF and Excel formats

This module provides reusable export functions for generating
downloadable PDF and Excel files from Django querysets.
"""

from django.http import HttpResponse
from django.utils import timezone
import io

# PDF generation
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_CENTER, TA_LEFT

# Excel generation
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment


def export_profiles_to_pdf(profiles_queryset, title="User Profiles Export", 
                           filters=None):
    """
    Export profiles to PDF format using reportlab.
    
    Args:
        profiles_queryset: QuerySet of Profile objects
        title: PDF document title (default: "User Profiles Export")
        filters: Optional dict with applied filters (e.g., {'date_from': '2025-01-01'})
    
    Returns:
        HttpResponse with PDF file as attachment
    
    Example:
        >>> from users.models import Profile
        >>> profiles = Profile.objects.all()
        >>> response = export_profiles_to_pdf(profiles)
        >>> return response  # Django view returns this
    """
    # Create buffer for PDF
    buffer = io.BytesIO()
    
    # Create PDF document
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )
    
    # Container for PDF elements
    elements = []
    
    # Define styles
    styles = getSampleStyleSheet()
    
    # Custom title style
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#1a237e'),
        spaceAfter=20,
        alignment=TA_CENTER,
        fontName='Helvetica-Bold'
    )
    
    # Custom subtitle style
    subtitle_style = ParagraphStyle(
        'CustomSubtitle',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor('#424242'),
        spaceAfter=5,
        alignment=TA_LEFT
    )
    
    # Add title
    title_paragraph = Paragraph(title, title_style)
    elements.append(title_paragraph)
    elements.append(Spacer(1, 0.2*inch))
    
    # Add metadata
    record_count = profiles_queryset.count()
    metadata_lines = [
        f"<b>Generated:</b> {timezone.now().strftime('%B %d, %Y at %H:%M:%S')}",
        f"<b>Total Profiles:</b> {record_count}",
    ]
    
    # Add filter information if provided
    if filters:
        if filters.get('date_from'):
            metadata_lines.append(f"<b>Date From:</b> {filters['date_from']}")
        if filters.get('date_to'):
            metadata_lines.append(f"<b>Date To:</b> {filters['date_to']}")
        if filters.get('gender'):
            metadata_lines.append(f"<b>Gender Filter:</b> {filters['gender'].title()}")
        if filters.get('location'):
            metadata_lines.append(f"<b>Location Filter:</b> {filters['location']}")
    
    for line in metadata_lines:
        elements.append(Paragraph(line, subtitle_style))
    
    elements.append(Spacer(1, 0.3*inch))
    
    # Prepare table data
    table_data = [
        ['#', 'Name', 'Gender', 'Age', 'Location', 'Occupation', 'Mobile Number']
    ]
    
    for idx, profile in enumerate(profiles_queryset, 1):
        table_data.append([
            str(idx),
            profile.name[:25] if len(profile.name) > 25 else profile.name,
            profile.gender.capitalize(),
            str(profile.age),
            profile.location[:30] if len(profile.location) > 30 else profile.location,
            profile.occupation[:20] if profile.occupation and len(profile.occupation) > 20 else (profile.occupation or 'N/A'),
            profile.mobile_number
        ])
    
    # Create table
    table = Table(table_data, repeatRows=1)
    
    # Apply table styling
    table.setStyle(TableStyle([
        # Header row styling
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a237e')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('TOPPADDING', (0, 0), (-1, 0), 12),
        
        # Data rows styling
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
        ('ALIGN', (0, 1), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        
        # Alternating row colors
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.beige, colors.lightgrey]),
        
        # Grid and borders
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('BOX', (0, 0), (-1, -1), 2, colors.HexColor('#1a237e')),
        
        # Padding
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 1), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 8),
    ]))
    
    elements.append(table)
    
    # Add footer
    elements.append(Spacer(1, 0.5*inch))
    footer_style = ParagraphStyle(
        'Footer',
        parent=styles['Normal'],
        fontSize=8,
        textColor=colors.grey,
        alignment=TA_CENTER
    )
    footer_text = Paragraph(
        "Generated by Matrimonial Website Admin Panel | Confidential Document",
        footer_style
    )
    elements.append(footer_text)
    
    # Build PDF
    doc.build(elements)
    
    # Get PDF data from buffer
    pdf_data = buffer.getvalue()
    buffer.close()
    
    # Create HTTP response
    response = HttpResponse(pdf_data, content_type='application/pdf')
    
    # Generate filename with timestamp
    filename = f"profiles_export_{timezone.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response['Content-Length'] = len(pdf_data)
    
    return response


def export_profiles_to_excel(profiles_queryset, sheet_name="Profiles", 
                             filters=None):
    """
    Export profiles to Excel format using pandas and openpyxl.
    
    Args:
        profiles_queryset: QuerySet of Profile objects
        sheet_name: Excel sheet name (default: "Profiles")
        filters: Optional dict with applied filters
    
    Returns:
        HttpResponse with Excel file as attachment
    
    Example:
        >>> from users.models import Profile
        >>> profiles = Profile.objects.filter(gender='female')
        >>> response = export_profiles_to_excel(profiles, sheet_name="Female Profiles")
        >>> return response  # Django view returns this
    """
    # Prepare data for DataFrame
    data_list = []
    
    for profile in profiles_queryset:
        data_list.append({
            'Profile ID': profile.id,
            'Name': profile.name,
            'Username': profile.user.username,
            'Email': profile.user.email,
            'Gender': profile.gender.capitalize(),
            'Age': profile.age,
            'Location': profile.location,
            'Occupation': profile.occupation if profile.occupation else 'N/A',
            'Education': profile.education if profile.education else 'N/A',
            'Height (ft)': float(profile.height) if profile.height else None,
            'Mobile Number': profile.mobile_number,
            'About': profile.about[:100] if profile.about else 'N/A',  # First 100 chars
            'Desired Partner Traits': profile.desired_partner_traits[:100] if profile.desired_partner_traits else 'N/A',
            'Profile Created': profile.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            'Last Updated': profile.updated_at.strftime('%Y-%m-%d %H:%M:%S'),
            'User Active': 'Yes' if profile.user.is_active else 'No',
        })
    
    # Create DataFrame
    df = pd.DataFrame(data_list)
    
    # Create Excel file in memory
    buffer = io.BytesIO()
    
    # Write to Excel with formatting
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        # Write main data
        df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=4)
        
        # Get workbook and worksheet
        workbook = writer.book
        worksheet = writer.sheets[sheet_name]
        
        # Add title and metadata at the top
        worksheet['A1'] = 'User Profiles Export Report'
        worksheet['A1'].font = Font(size=16, bold=True, color='1a237e')
        
        worksheet['A2'] = f"Generated: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}"
        worksheet['A2'].font = Font(size=10, italic=True)
        
        worksheet['A3'] = f"Total Records: {len(df)}"
        worksheet['A3'].font = Font(size=10, italic=True)
        
        # Add filter information if provided
        if filters:
            filter_text = "Filters Applied: "
            filter_parts = []
            if filters.get('date_from'):
                filter_parts.append(f"Date From: {filters['date_from']}")
            if filters.get('date_to'):
                filter_parts.append(f"Date To: {filters['date_to']}")
            if filters.get('gender'):
                filter_parts.append(f"Gender: {filters['gender']}")
            if filters.get('location'):
                filter_parts.append(f"Location: {filters['location']}")
            
            if filter_parts:
                worksheet['A4'] = filter_text + ", ".join(filter_parts)
                worksheet['A4'].font = Font(size=9, italic=True, color='666666')
        
        # Style header row (row 5, after metadata)
        header_fill = PatternFill(start_color='1a237e', end_color='1a237e', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        for cell in worksheet[5]:  # Row 5 is the header (startrow=4, 0-indexed)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        # Auto-adjust column widths
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if cell.value:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length:
                            max_length = cell_length
                except:
                    pass
            
            # Set adjusted width (min 10, max 50)
            adjusted_width = min(max(max_length + 2, 10), 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
        
        # Format data cells
        data_alignment = Alignment(vertical='center', wrap_text=True)
        for row in worksheet.iter_rows(min_row=6, max_row=len(df)+5):  # Data rows
            for cell in row:
                cell.alignment = data_alignment
        
        # Freeze header row
        worksheet.freeze_panes = 'A6'  # Freeze rows above row 6
        
        # Add alternating row colors
        light_fill = PatternFill(start_color='F5F5F5', end_color='F5F5F5', fill_type='solid')
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=6, max_row=len(df)+5), start=1):
            if row_idx % 2 == 0:  # Even rows
                for cell in row:
                    cell.fill = light_fill
    
    # Get Excel data from buffer
    excel_data = buffer.getvalue()
    buffer.close()
    
    # Create HTTP response
    response = HttpResponse(
        excel_data,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    
    # Generate filename with timestamp
    filename = f"profiles_export_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response['Content-Length'] = len(excel_data)
    
    return response


def export_profiles_to_csv(profiles_queryset, filename_prefix="profiles"):
    """
    Export profiles to CSV format using pandas.
    
    Args:
        profiles_queryset: QuerySet of Profile objects
        filename_prefix: Prefix for the CSV filename
    
    Returns:
        HttpResponse with CSV file as attachment
    
    Example:
        >>> from users.models import Profile
        >>> profiles = Profile.objects.all()
        >>> response = export_profiles_to_csv(profiles)
        >>> return response
    """
    # Prepare data for DataFrame
    data_list = []
    
    for profile in profiles_queryset:
        data_list.append({
            'Profile ID': profile.id,
            'Name': profile.name,
            'Username': profile.user.username,
            'Email': profile.user.email,
            'Gender': profile.gender,
            'Age': profile.age,
            'Location': profile.location,
            'Occupation': profile.occupation if profile.occupation else '',
            'Education': profile.education if profile.education else '',
            'Height': float(profile.height) if profile.height else '',
            'Mobile Number': profile.mobile_number,
            'Created': profile.created_at.strftime('%Y-%m-%d %H:%M:%S'),
        })
    
    # Create DataFrame
    df = pd.DataFrame(data_list)
    
    # Create CSV in memory
    buffer = io.StringIO()
    df.to_csv(buffer, index=False, encoding='utf-8')
    
    # Get CSV data
    csv_data = buffer.getvalue()
    buffer.close()
    
    # Create HTTP response
    response = HttpResponse(csv_data, content_type='text/csv')
    
    # Generate filename with timestamp
    filename = f"{filename_prefix}_export_{timezone.now().strftime('%Y%m%d_%H%M%S')}.csv"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    return response
